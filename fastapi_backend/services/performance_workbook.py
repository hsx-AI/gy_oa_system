# -*- coding: utf-8 -*-
"""Performance workbook builder for the pilot department."""
from __future__ import annotations

import hashlib
import logging
import re
import shutil
from copy import copy
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

logger = logging.getLogger(__name__)

PILOT_DEPARTMENT = "智能制造技术室"
LOCK_PASSWORD = "perf-lock"
_CELL_REF = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")
_INVALID_SHEET = re.compile(r"[\[\]\:\*\?\/\\]")


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def safe_sheet_title(name: str) -> str:
    title = _INVALID_SHEET.sub("_", (name or "").strip()) or "会议纪实"
    return title[:31]


def find_layout(ws: Worksheet) -> Tuple[int, int, int]:
    """Return name column, header end row, and sequence column."""
    name_col = 2
    header_end = 1
    seq_col = 1
    for row in range(1, min(8, ws.max_row or 1) + 1):
        for col in range(1, min(8, ws.max_column or 1) + 1):
            value = ws.cell(row, col).value
            if not isinstance(value, str):
                continue
            text = value.strip()
            if text == "姓名":
                name_col = col
                header_end = row
            elif text == "序号":
                seq_col = col
                header_end = max(header_end, row)
    return name_col, header_end, seq_col


def people_rows(ws: Worksheet, name_col: int, header_end: int) -> List[Tuple[int, str]]:
    found: List[Tuple[int, str]] = []
    blanks = 0
    for row in range(header_end + 1, (ws.max_row or header_end) + 1):
        value = ws.cell(row, name_col).value
        name = value.strip() if isinstance(value, str) else ""
        if name:
            found.append((row, name))
            blanks = 0
            continue
        blanks += 1
        if found and blanks >= 5:
            break
    return found


def formula_columns(ws: Worksheet, row: int) -> set:
    cols = set()
    for col in range(1, (ws.max_column or 1) + 1):
        value = ws.cell(row, col).value
        if isinstance(value, str) and value.startswith("="):
            cols.add(col)
    return cols


def _shift_formula(formula: str, src_row: int, dst_row: int) -> str:
    def repl(match: re.Match) -> str:
        col_abs, col, row_abs, row_text = match.groups()
        if int(row_text) == src_row:
            return f"{col_abs}{col}{row_abs}{dst_row}"
        return match.group(0)

    return _CELL_REF.sub(repl, formula)


def _copy_row(ws: Worksheet, src_row: int, dst_row: int) -> None:
    if src_row == dst_row:
        return
    max_col = ws.max_column or 1
    for col in range(1, max_col + 1):
        src = ws.cell(src_row, col)
        if isinstance(src, MergedCell):
            continue
        dst = ws.cell(dst_row, col)
        if isinstance(dst, MergedCell):
            continue
        value = src.value
        if isinstance(value, str) and value.startswith("="):
            dst.value = _shift_formula(value, src_row, dst_row)
        else:
            dst.value = value
        if src.has_style:
            dst.font = copy(src.font)
            dst.border = copy(src.border)
            dst.fill = copy(src.fill)
            dst.number_format = src.number_format
            dst.alignment = copy(src.alignment)
            dst.protection = copy(src.protection)
    src_dim = ws.row_dimensions[src_row]
    if src_dim.height:
        ws.row_dimensions[dst_row].height = src_dim.height


def _read_inputs(ws: Worksheet, row: int, skip_cols: Iterable[int], include_blank: bool = False) -> Dict[int, object]:
    skipped = set(skip_cols)
    inputs: Dict[int, object] = {}
    for col in range(1, (ws.max_column or 1) + 1):
        if col in skipped:
            continue
        cell = ws.cell(row, col)
        if isinstance(cell, MergedCell):
            continue
        value = cell.value
        if isinstance(value, str) and value.startswith("="):
            continue
        if value is None or value == "":
            if include_blank:
                inputs[col] = None
            continue
        inputs[col] = value
    return inputs


def _write_inputs(ws: Worksheet, row: int, inputs: Dict[int, object], formula_cols: Iterable[int]) -> None:
    blocked = set(formula_cols)
    for col, value in inputs.items():
        if col in blocked:
            continue
        cell = ws.cell(row, col)
        if isinstance(cell, MergedCell):
            continue
        current = cell.value
        if isinstance(current, str) and current.startswith("="):
            continue
        cell.value = value


def _meeting_values(ws: Worksheet) -> Dict[Tuple[int, int], object]:
    values: Dict[Tuple[int, int], object] = {}
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            value = cell.value
            if value is None or value == "":
                continue
            if isinstance(value, str) and value.startswith("="):
                continue
            if cell.row <= 3 and cell.column == 1:
                continue
            values[(cell.row, cell.column)] = value
    return values


def _apply_meeting_values(ws: Worksheet, values: Dict[Tuple[int, int], object]) -> None:
    for (row, col), value in values.items():
        cell = ws.cell(row, col)
        if isinstance(cell, MergedCell):
            continue
        current = cell.value
        if isinstance(current, str) and current.startswith("="):
            continue
        cell.value = value


def _blank_meeting_body(ws: Worksheet) -> None:
    for row in ws.iter_rows(min_row=4, min_col=2):
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            if isinstance(cell.value, str) and str(cell.value).startswith("="):
                continue
            cell.value = None


def clone_sheet(src: Worksheet, dest_wb: Workbook, title: str) -> Worksheet:
    title = safe_sheet_title(title)
    if title in dest_wb.sheetnames:
        dest_wb.remove(dest_wb[title])
    ws = dest_wb.create_sheet(title)
    for letter, dim in src.column_dimensions.items():
        ws.column_dimensions[letter].width = dim.width
        ws.column_dimensions[letter].hidden = dim.hidden
    for index, dim in src.row_dimensions.items():
        if dim.height is not None:
            ws.row_dimensions[index].height = dim.height
        ws.row_dimensions[index].hidden = dim.hidden
    for merged in list(src.merged_cells.ranges):
        ws.merge_cells(str(merged))
    for row in src.iter_rows():
        for cell in row:
            if isinstance(cell, MergedCell):
                continue
            target = ws.cell(cell.row, cell.column)
            if isinstance(target, MergedCell):
                continue
            target.value = cell.value
            if cell.has_style:
                target.font = copy(cell.font)
                target.border = copy(cell.border)
                target.fill = copy(cell.fill)
                target.number_format = cell.number_format
                target.alignment = copy(cell.alignment)
                target.protection = copy(cell.protection)
    ws.freeze_panes = src.freeze_panes
    ws.sheet_view.showGridLines = src.sheet_view.showGridLines
    ws.page_setup.orientation = src.page_setup.orientation
    ws.page_setup.paperSize = src.page_setup.paperSize
    ws.page_setup.fitToPage = src.page_setup.fitToPage
    ws.page_setup.fitToWidth = src.page_setup.fitToWidth
    ws.page_setup.fitToHeight = src.page_setup.fitToHeight
    ws.sheet_properties.pageSetUpPr.fitToPage = src.sheet_properties.pageSetUpPr.fitToPage
    if src.oddHeader.left.text or src.oddHeader.center.text or src.oddHeader.right.text:
        ws.oddHeader.left.text = src.oddHeader.left.text
        ws.oddHeader.center.text = src.oddHeader.center.text
        ws.oddHeader.right.text = src.oddHeader.right.text
    return ws


def _unlock_sheet(ws: Worksheet) -> None:
    ws.protection.sheet = False


def _protect_score_sheet(ws: Worksheet, editable_row: Optional[int], formula_cols: Iterable[int]) -> None:
    blocked = set(formula_cols)
    if editable_row:
        for col in range(1, (ws.max_column or 1) + 1):
            if col in blocked or col <= 2:
                continue
            cell = ws.cell(editable_row, col)
            if isinstance(cell, MergedCell):
                continue
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                continue
            cell.protection = Protection(locked=False)
    ws.protection.sheet = True
    ws.protection.password = LOCK_PASSWORD
    ws.protection.insertRows = False
    ws.protection.insertColumns = False
    ws.protection.deleteRows = False
    ws.protection.deleteColumns = False


def _find_total_col(ws: Worksheet, header_end: int) -> Optional[int]:
    for row in range(1, header_end + 1):
        for col in range(1, (ws.max_column or 1) + 1):
            value = ws.cell(row, col).value
            if isinstance(value, str) and value.strip() == "合计":
                return col
    return None


def _add_rank_column(ws: Worksheet, header_end: int, data_rows: Sequence[int]) -> None:
    total_col = _find_total_col(ws, header_end)
    if not total_col or not data_rows:
        return
    rank_col = total_col + 1
    head = ws.cell(1, rank_col)
    if isinstance(head, MergedCell) or (head.value not in (None, "")):
        return
    letter = get_column_letter(total_col)
    first, last = data_rows[0], data_rows[-1]
    if header_end > 1:
        ws.merge_cells(start_row=1, start_column=rank_col, end_row=header_end, end_column=rank_col)
    cell = ws.cell(1, rank_col, "排名")
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1D4ED8")
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in data_rows:
        ws.cell(
            row,
            rank_col,
            f'=IF({letter}{row}="","",RANK({letter}{row},{letter}${first}:{letter}${last},0))',
        )
    ws.column_dimensions[get_column_letter(rank_col)].width = 10


def _isolate_person_row(ws: Worksheet, employee: str) -> Tuple[int, int]:
    name_col, header_end, seq_col = find_layout(ws)
    rows = people_rows(ws, name_col, header_end)
    if not rows:
        raise ValueError("template score sheet has no employee rows")
    first_row = rows[0][0]
    matched = next((row for row, name in rows if name == employee), None)
    if matched and matched != first_row:
        _copy_row(ws, matched, first_row)
    elif matched is None:
        for col in range(1, (ws.max_column or 1) + 1):
            cell = ws.cell(first_row, col)
            if isinstance(cell, MergedCell):
                continue
            if isinstance(cell.value, str) and str(cell.value).startswith("="):
                continue
            cell.value = None
    ws.cell(first_row, name_col).value = employee
    seq_cell = ws.cell(first_row, seq_col)
    if not isinstance(seq_cell, MergedCell):
        seq_cell.value = 1
    extra = (ws.max_row or first_row) - first_row
    if extra > 0:
        ws.delete_rows(first_row + 1, extra)
    return first_row, header_end


def _load_personal_payload(path: Optional[Path], employee: str) -> Tuple[Dict[int, object], Dict[Tuple[int, int], object]]:
    if not path or not path.is_file():
        return {}, {}
    wb = load_workbook(path)
    try:
        score = wb.worksheets[0]
        name_col, header_end, seq_col = find_layout(score)
        row_no = header_end + 1
        for row, name in people_rows(score, name_col, header_end):
            if name == employee:
                row_no = row
                break
        formulas = formula_columns(score, row_no)
        inputs = _read_inputs(score, row_no, set(formulas) | {name_col, seq_col})
        meeting: Dict[Tuple[int, int], object] = {}
        title = safe_sheet_title(employee)
        if title in wb.sheetnames:
            meeting = _meeting_values(wb[title])
        elif len(wb.worksheets) > 1:
            meeting = _meeting_values(wb.worksheets[1])
        return inputs, meeting
    finally:
        wb.close()


def validate_template(path: Path) -> None:
    wb = load_workbook(path)
    try:
        if len(wb.worksheets) < 2:
            raise ValueError("template needs a score sheet and a meeting sheet")
        ws = wb.worksheets[0]
        name_col, header_end, _seq = find_layout(ws)
        if header_end < 1:
            raise ValueError("score sheet header was not found")
        header_hit = False
        for row in range(1, header_end + 1):
            value = ws.cell(row, name_col).value
            if isinstance(value, str) and value.strip() == "姓名":
                header_hit = True
        if not header_hit:
            raise ValueError("score sheet header must contain the name column")
        if not people_rows(ws, name_col, header_end):
            raise ValueError("score sheet has no employee names")
    finally:
        wb.close()


def build_personal_workbook(
    *,
    template_path: Path,
    dest_path: Path,
    employee: str,
    previous_path: Optional[Path] = None,
) -> None:
    inputs, meeting_values = _load_personal_payload(previous_path, employee)
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template_path, dest_path)
    wb = load_workbook(dest_path)
    try:
        while len(wb.worksheets) > 2:
            wb.remove(wb.worksheets[-1])
        score = wb.worksheets[0]
        data_row, _header_end = _isolate_person_row(score, employee)
        name_col, header_end, seq_col = find_layout(score)
        formulas = formula_columns(score, data_row)
        _write_inputs(score, data_row, inputs, formulas | {name_col, seq_col})
        _protect_score_sheet(score, data_row, formulas | {1, 2})
        meeting = wb.worksheets[1]
        meeting.title = safe_sheet_title(employee)
        _blank_meeting_body(meeting)
        _apply_meeting_values(meeting, meeting_values)
        _unlock_sheet(meeting)
        wb.save(dest_path)
    finally:
        wb.close()


def personal_fill_state(path: Path, employee: str) -> bool:
    if not path.is_file():
        return False
    wb = load_workbook(path, data_only=False)
    try:
        score = wb.worksheets[0]
        name_col, header_end, seq_col = find_layout(score)
        for row in score.iter_rows(min_row=header_end + 1, max_row=header_end + 3):
            for cell in row:
                if cell.column in (name_col, seq_col):
                    continue
                value = cell.value
                if value is None or value == "":
                    continue
                if isinstance(value, str) and value.startswith("="):
                    continue
                return True
        title = safe_sheet_title(employee)
        meeting = wb[title] if title in wb.sheetnames else (wb.worksheets[1] if len(wb.worksheets) > 1 else None)
        if meeting is None:
            return False
        for row in meeting.iter_rows(min_row=4, min_col=2, max_row=min(meeting.max_row or 4, 30), max_col=min(meeting.max_column or 2, 8)):
            for cell in row:
                if cell.value not in (None, ""):
                    return True
        return False
    finally:
        wb.close()


def build_summary_workbook(
    *,
    template_path: Path,
    dest_path: Path,
    owner: str,
    roster: Sequence[str],
    personal_paths: Dict[str, Path],
) -> None:
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copyfile(template_path, dest_path)
    proto = load_workbook(template_path)
    wb = load_workbook(dest_path)
    try:
        score = wb.worksheets[0]
        name_col, header_end, seq_col = find_layout(score)
        existing = people_rows(score, name_col, header_end)
        if not existing:
            raise ValueError("template has no employee rows")
        first_row = existing[0][0]
        by_name = {name: row for row, name in existing}
        ordered = [name for _row, name in existing]
        for name in roster:
            if name and name not in by_name:
                ordered.append(name)
        next_row = existing[-1][0]
        for name in ordered:
            if name in by_name:
                continue
            next_row += 1
            _copy_row(score, first_row, next_row)
            score.cell(next_row, name_col).value = name
            score.cell(next_row, seq_col).value = len(by_name) + 1
            by_name[name] = next_row
        formula_cols = formula_columns(score, first_row)
        for name, row in by_name.items():
            personal = personal_paths.get(name)
            inputs, _meeting = _load_personal_payload(personal, name)
            _write_inputs(score, row, inputs, formula_cols | {name_col, seq_col})
        data_row_numbers = [by_name[name] for name in ordered if name in by_name]
        _add_rank_column(score, header_end, data_row_numbers)
        owner_row = by_name.get(owner)
        _protect_score_sheet(score, owner_row, formula_cols | {1, 2})

        proto_meeting = proto.worksheets[1]
        for sheet in list(wb.worksheets[1:]):
            wb.remove(sheet)
        for name in ordered:
            personal = personal_paths.get(name)
            source = None
            if personal and personal.is_file():
                person_wb = load_workbook(personal)
                try:
                    title = safe_sheet_title(name)
                    if title in person_wb.sheetnames:
                        source_ws = person_wb[title]
                    elif len(person_wb.worksheets) > 1:
                        source_ws = person_wb.worksheets[1]
                    else:
                        source_ws = None
                    if source_ws is not None:
                        cloned = clone_sheet(source_ws, wb, name)
                        source = cloned
                finally:
                    person_wb.close()
            if source is None:
                source = clone_sheet(proto_meeting, wb, name)
                _blank_meeting_body(source)
            if name != owner:
                source.protection.sheet = True
                source.protection.password = LOCK_PASSWORD
            else:
                _unlock_sheet(source)
        wb.save(dest_path)
    finally:
        wb.close()
        proto.close()


def sync_owner_personal_from_summary(
    *,
    summary_path: Path,
    template_path: Path,
    personal_path: Path,
    owner: str,
) -> None:
    """Copy the owner row and meeting sheet from a summary save into the personal file."""
    saved = load_workbook(summary_path)
    try:
        score = saved.worksheets[0]
        name_col, header_end, seq_col = find_layout(score)
        owner_row = None
        for row, name in people_rows(score, name_col, header_end):
            if name == owner:
                owner_row = row
                break
        if owner_row is None:
            raise ValueError("owner row was not found in the summary workbook: " + owner)
        template_wb = load_workbook(template_path)
        try:
            template_score = template_wb.worksheets[0]
            t_name, t_header, t_seq = find_layout(template_score)
            t_rows = people_rows(template_score, t_name, t_header)
            formula_row = t_rows[0][0] if t_rows else t_header + 1
            formulas = formula_columns(template_score, formula_row)
        finally:
            template_wb.close()
        inputs = _read_inputs(score, owner_row, formulas | {name_col, seq_col, 1, 2}, include_blank=True)
        title = safe_sheet_title(owner)
        meeting_ws = saved[title] if title in saved.sheetnames else None
        meeting_values = _meeting_values(meeting_ws) if meeting_ws is not None else {}
    finally:
        saved.close()

    build_personal_workbook(
        template_path=template_path,
        dest_path=personal_path,
        employee=owner,
        previous_path=None,
    )
    wb = load_workbook(personal_path)
    try:
        personal_score = wb.worksheets[0]
        p_name, p_header, p_seq = find_layout(personal_score)
        data_row = p_header + 1
        for row, name in people_rows(personal_score, p_name, p_header):
            if name == owner:
                data_row = row
                break
        formulas = formula_columns(personal_score, data_row)
        _write_inputs(personal_score, data_row, inputs, formulas | {p_name, p_seq})
        if len(wb.worksheets) > 1:
            _apply_meeting_values(wb.worksheets[1], meeting_values)
        wb.save(personal_path)
    finally:
        wb.close()
