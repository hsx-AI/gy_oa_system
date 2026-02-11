# -*- coding: utf-8 -*-
"""
系统管理员 - 数据库表增删改查
仅 webconfig.admin1 对应用户（与 yggl.name 映射）可访问。
"""
import re
import logging
from typing import Optional, List, Dict, Any

import os
import tempfile
from fastapi import APIRouter, HTTPException, Query, File, UploadFile, Form
from pydantic import BaseModel
from database import db

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/db-manager", tags=["系统管理员-数据库"])

# 表名/列名只允许字母数字下划线，防止 SQL 注入
_NAME_PATTERN = re.compile(r"^[a-zA-Z0-9_]+$")


def _get_admin1() -> Optional[str]:
    """从 webconfig 表读取 admin1（系统管理员用户名，对应 yggl.name）。"""
    try:
        rows = db.execute_query(
            "SELECT admin1 FROM webconfig WHERE id = %s LIMIT 1",
            ("1",),
        )
        if rows and rows[0].get("admin1") is not None:
            return (rows[0]["admin1"] or "").strip() or None
    except Exception as e:
        logger.debug(f"读取 webconfig.admin1 失败（可能无此列）: {e}")
    return None


def _require_system_admin(current_user: str) -> None:
    """校验当前用户是否为系统管理员（admin1），否则 403。"""
    admin1 = _get_admin1()
    if not admin1 or (current_user or "").strip() != admin1:
        raise HTTPException(status_code=403, detail="仅系统管理员（webconfig.admin1）可操作")


def _validate_identifier(name: str, kind: str = "表名") -> None:
    if not name or not _NAME_PATTERN.match(name):
        raise HTTPException(status_code=400, detail=f"无效的{kind}")


@router.get("/permission")
async def db_manager_permission(
    current_user: str = Query(..., description="当前登录用户姓名"),
):
    """
    检查当前用户是否有数据库管理权限（webconfig.admin1）。
    返回: { canAccess: true/false }
    """
    admin1 = _get_admin1()
    can = bool(admin1 and (current_user or "").strip() == admin1)
    return {"success": True, "canAccess": can}


@router.get("/tables")
async def list_tables(
    current_user: str = Query(..., description="当前登录用户，用于权限校验"),
):
    """列出当前数据库下所有表名。仅系统管理员可访问。"""
    _require_system_admin(current_user)
    try:
        rows = db.execute_query(
            "SELECT TABLE_NAME AS name FROM information_schema.TABLES "
            "WHERE TABLE_SCHEMA = DATABASE() ORDER BY TABLE_NAME"
        )
        names = [r["name"] for r in (rows or []) if r.get("name")]
        return {"success": True, "list": names}
    except Exception as e:
        logger.error(f"列出表失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/table/{table_name}/columns")
async def get_table_columns(
    table_name: str,
    current_user: str = Query(..., description="当前登录用户，用于权限校验"),
):
    """获取指定表的列信息（含主键）。仅系统管理员可访问。"""
    _require_system_admin(current_user)
    _validate_identifier(table_name)

    try:
        cols = db.execute_query(
            "SELECT COLUMN_NAME AS name, DATA_TYPE AS type, IS_NULLABLE AS nullable, "
            "COLUMN_KEY AS `key`, COLUMN_DEFAULT AS `default` "
            "FROM information_schema.COLUMNS "
            "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s ORDER BY ORDINAL_POSITION",
            (table_name,),
        )
        if not cols:
            raise HTTPException(status_code=404, detail="表不存在或无列")
        pk_rows = db.execute_query(
            "SELECT COLUMN_NAME AS name FROM information_schema.KEY_COLUMN_USAGE "
            "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s AND CONSTRAINT_NAME = 'PRIMARY' "
            "ORDER BY ORDINAL_POSITION",
            (table_name,),
        )
        pk_set = {r["name"] for r in (pk_rows or []) if r.get("name")}
        list_data = []
        for c in cols:
            list_data.append({
                "name": c.get("name"),
                "type": c.get("type"),
                "nullable": (c.get("nullable") or "").upper() == "YES",
                "key": c.get("key") or "",
                "default": c.get("default"),
                "isPk": c.get("name") in pk_set,
            })
        return {"success": True, "list": list_data, "primaryKey": list(pk_set)}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"获取表列失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/table/{table_name}/rows")
async def get_table_rows(
    table_name: str,
    current_user: str = Query(..., description="当前登录用户，用于权限校验"),
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
    search_column: Optional[str] = Query(None, description="要搜索的列名"),
    search_keyword: Optional[str] = Query(None, description="搜索关键词（模糊匹配）"),
):
    """分页查询指定表数据。仅系统管理员可访问。支持按列模糊搜索。"""
    _require_system_admin(current_user)
    _validate_identifier(table_name)

    try:
        safe_table = f"`{table_name}`"
        where_clause = ""
        count_params: list = []
        select_params: list = []

        if search_column and search_keyword is not None and (search_keyword.strip() or ""):
            _validate_identifier(search_column, "列名")
            cols = db.execute_query(
                "SELECT COLUMN_NAME AS name FROM information_schema.COLUMNS "
                "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s",
                (table_name,),
            )
            col_names = {c["name"] for c in (cols or []) if c.get("name")}
            if search_column not in col_names:
                raise HTTPException(status_code=400, detail=f"列 {search_column} 不存在于该表")
            where_clause = f" WHERE `{search_column}` LIKE %s"
            keyword_param = f"%{search_keyword.strip()}%"
            count_params.append(keyword_param)
            select_params.append(keyword_param)

        count_sql = f"SELECT COUNT(*) AS cnt FROM {safe_table}{where_clause}"
        total = db.execute_scalar(count_sql, tuple(count_params)) if count_params else db.execute_scalar(count_sql)
        total = total or 0

        offset = (page - 1) * page_size
        select_sql = f"SELECT * FROM {safe_table}{where_clause} LIMIT %s OFFSET %s"
        select_params.extend([page_size, offset])
        rows = db.execute_query(select_sql, tuple(select_params))

        list_data = []
        for r in rows or []:
            row_dict = {}
            for k, v in r.items():
                if hasattr(v, "isoformat"):
                    row_dict[k] = v.isoformat() if v else None
                else:
                    row_dict[k] = v
            list_data.append(row_dict)
        return {"success": True, "list": list_data, "total": total, "page": page, "pageSize": page_size}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"查询表数据失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))


class InsertRowRequest(BaseModel):
    current_user: str
    row: Dict[str, Any]  # 列名 -> 值，None 表示 NULL


@router.post("/table/{table_name}/rows")
async def insert_table_row(
    table_name: str,
    req: InsertRowRequest,
):
    """向指定表插入一行。仅系统管理员可访问。"""
    _require_system_admin(req.current_user)
    _validate_identifier(table_name)
    if not req.row:
        raise HTTPException(status_code=400, detail="row 不能为空")

    try:
        cols = db.execute_query(
            "SELECT COLUMN_NAME AS name FROM information_schema.COLUMNS "
            "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s ORDER BY ORDINAL_POSITION",
            (table_name,),
        )
        col_names = [c["name"] for c in (cols or []) if c.get("name")]
        if not col_names:
            raise HTTPException(status_code=404, detail="表不存在或无列")
        # 只使用请求中且表中存在的列
        valid = {k: v for k, v in req.row.items() if _NAME_PATTERN.match(str(k)) and k in col_names}
        if not valid:
            raise HTTPException(status_code=400, detail="未提供有效列")
        placeholders = ", ".join(["%s"] * len(valid))
        columns = ", ".join(f"`{k}`" for k in valid.keys())
        safe_table = f"`{table_name}`"
        sql = f"INSERT INTO {safe_table} ({columns}) VALUES ({placeholders})"
        params = tuple(valid.values())
        db.execute_update(sql, params)
        return {"success": True, "message": "插入成功"}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"插入失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))


class UpdateRowRequest(BaseModel):
    current_user: str
    row: Dict[str, Any]  # 必须包含主键列，其余为要更新的列


@router.put("/table/{table_name}/rows")
async def update_table_row(
    table_name: str,
    req: UpdateRowRequest,
):
    """更新指定表一行（按主键定位）。仅系统管理员可访问。"""
    _require_system_admin(req.current_user)
    _validate_identifier(table_name)
    if not req.row:
        raise HTTPException(status_code=400, detail="row 不能为空")

    try:
        pk_rows = db.execute_query(
            "SELECT COLUMN_NAME AS name FROM information_schema.KEY_COLUMN_USAGE "
            "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s AND CONSTRAINT_NAME = 'PRIMARY' "
            "ORDER BY ORDINAL_POSITION",
            (table_name,),
        )
        pk_cols = [r["name"] for r in (pk_rows or []) if r.get("name")]
        if not pk_cols:
            raise HTTPException(status_code=400, detail="该表无主键，无法按主键更新")
        cols = db.execute_query(
            "SELECT COLUMN_NAME AS name FROM information_schema.COLUMNS "
            "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s",
            (table_name,),
        )
        all_cols = {c["name"] for c in (cols or []) if c.get("name")}
        for pk in pk_cols:
            if pk not in req.row:
                raise HTTPException(status_code=400, detail=f"缺少主键字段: {pk}")
        valid = {k: v for k, v in req.row.items() if _NAME_PATTERN.match(str(k)) and k in all_cols}
        if len(valid) <= len(pk_cols):
            raise HTTPException(status_code=400, detail="请提供除主键外至少一个要更新的列")
        set_parts = []
        set_params = []
        for k, v in valid.items():
            if k in pk_cols:
                continue
            set_parts.append(f"`{k}` = %s")
            set_params.append(v)
        where_parts = [f"`{k}` = %s" for k in pk_cols]
        set_params.extend([req.row[k] for k in pk_cols])
        safe_table = f"`{table_name}`"
        sql = f"UPDATE {safe_table} SET {', '.join(set_parts)} WHERE {' AND '.join(where_parts)}"
        n = db.execute_update(sql, tuple(set_params))
        return {"success": True, "message": "更新成功", "affected": n}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"更新失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))


class DeleteRowRequest(BaseModel):
    current_user: str
    row: Dict[str, Any]  # 至少包含主键列，用于 WHERE 条件


@router.delete("/table/{table_name}/rows")
async def delete_table_row(
    table_name: str,
    req: DeleteRowRequest,
):
    """按主键删除指定表一行。仅系统管理员可访问。"""
    _require_system_admin(req.current_user)
    _validate_identifier(table_name)
    if not req.row:
        raise HTTPException(status_code=400, detail="row 不能为空")

    try:
        pk_rows = db.execute_query(
            "SELECT COLUMN_NAME AS name FROM information_schema.KEY_COLUMN_USAGE "
            "WHERE TABLE_SCHEMA = DATABASE() AND TABLE_NAME = %s AND CONSTRAINT_NAME = 'PRIMARY' "
            "ORDER BY ORDINAL_POSITION",
            (table_name,),
        )
        pk_cols = [r["name"] for r in (pk_rows or []) if r.get("name")]
        if not pk_cols:
            raise HTTPException(status_code=400, detail="该表无主键，无法按主键删除")
        for pk in pk_cols:
            if pk not in req.row:
                raise HTTPException(status_code=400, detail=f"缺少主键字段: {pk}")
        where_parts = [f"`{k}` = %s" for k in pk_cols]
        params = [req.row[k] for k in pk_cols]
        safe_table = f"`{table_name}`"
        sql = f"DELETE FROM {safe_table} WHERE {' AND '.join(where_parts)}"
        n = db.execute_update(sql, tuple(params))
        return {"success": True, "message": "删除成功", "affected": n}
    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"删除失败: {e}")
        raise HTTPException(status_code=500, detail=str(e))


# ---------- yggl 按身份证号批量填充（系统管理员） ----------
# 允许通过 Excel（A列身份证号、B列填充值）批量更新 yggl 的指定字段
YGGL_FILL_FIELDS = {
    "gh": "工号",
    "lsys": "所在科室",
    "jb": "级别",
    "rcnf": "入厂时间",
    "sfzh": "身份证号",
    "xbie": "性别",
    "name": "姓名",
}


def _read_excel_ab(path: str) -> List[tuple]:
    """读取 Excel 前两列，返回 [(A, B), ...]。支持 .xlsx / .xls。"""
    path_lower = path.lower()
    if path_lower.endswith(".xlsx"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            ws = wb.active
            rows = []
            for row in ws.iter_rows(min_row=1, max_col=2, values_only=True):
                a = row[0] if len(row) > 0 else None
                b = row[1] if len(row) > 1 else None
                if a is not None or b is not None:
                    rows.append((a, b))
            wb.close()
            return rows
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"读取 xlsx 失败: {e}")
    if path_lower.endswith(".xls"):
        try:
            import xlrd
            with xlrd.open_workbook(path) as book:
                sheet = book.sheet_by_index(0)
                return [
                    (
                        sheet.cell_value(i, 0) if sheet.ncols > 0 else None,
                        sheet.cell_value(i, 1) if sheet.ncols > 1 else None,
                    )
                    for i in range(sheet.nrows)
                ]
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"读取 xls 失败: {e}")
    raise HTTPException(status_code=400, detail="仅支持 .xlsx 或 .xls 文件")


@router.get("/yggl-fill-fields")
async def get_yggl_fill_fields(
    current_user: str = Query(..., description="当前登录用户，用于权限校验"),
):
    """返回可批量填充的 yggl 字段列表（仅系统管理员）。"""
    _require_system_admin(current_user)
    return {
        "success": True,
        "list": [{"value": k, "label": v} for k, v in YGGL_FILL_FIELDS.items()],
    }


@router.post("/yggl-fill-by-excel")
async def yggl_fill_by_excel(
    current_user: str = Query(..., description="当前登录用户，用于权限校验"),
    field: str = Form(..., description="要填充的 yggl 字段名"),
    file: UploadFile = File(..., description="Excel：A列身份证号，B列填充值"),
):
    """
    按 Excel 批量更新 yggl 指定字段。仅系统管理员可操作。
    Excel 格式：A列 = 身份证号（用于匹配），B列 = 要写入的值。
    表头行可自动识别并跳过。
    """
    _require_system_admin(current_user)
    if field not in YGGL_FILL_FIELDS:
        raise HTTPException(status_code=400, detail=f"无效字段，可选: {list(YGGL_FILL_FIELDS.keys())}")
    if not file.filename:
        raise HTTPException(status_code=400, detail="请选择文件")
    suffix = os.path.splitext(file.filename)[1].lower()
    if suffix not in (".xlsx", ".xls"):
        raise HTTPException(status_code=400, detail="仅支持 .xlsx 或 .xls 文件")
    try:
        content = await file.read()
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"读取文件失败: {e}")
    tmp = None
    try:
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as f:
            f.write(content)
            tmp = f.name
        rows = _read_excel_ab(tmp)
    finally:
        if tmp and os.path.isfile(tmp):
            try:
                os.unlink(tmp)
            except Exception:
                pass
    if not rows:
        return {"success": True, "updated": 0, "unmapped": [], "message": "表格无数据"}
    # 跳过表头
    first_a = str(rows[0][0] or "").strip()
    first_b = str(rows[0][1] or "").strip()
    if first_a in ("身份证号", "身份证", "sfzh") or first_b in ("填充值", "值", "value"):
        rows = rows[1:]
    if not rows:
        return {"success": True, "updated": 0, "unmapped": [], "message": "除表头外无数据"}
    # 规范化：身份证号去空格
    pairs = []
    for a, b in rows:
        sfzh = str(a or "").strip().replace(" ", "")
        if not sfzh:
            continue
        val = b if b is None else str(b).strip()
        pairs.append((sfzh, val))
    if not pairs:
        return {"success": True, "updated": 0, "unmapped": [], "message": "没有有效的身份证号列"}
    updated = 0
    unmapped = []
    # 使用参数化：列名来自白名单，安全
    col = field
    for sfzh, val in pairs:
        exist = db.execute_query(
            "SELECT name FROM yggl WHERE REPLACE(TRIM(COALESCE(sfzh,'')), ' ', '') = %s LIMIT 1",
            (sfzh,),
        )
        if not exist:
            unmapped.append(sfzh)
            continue
        name = (exist[0].get("name") or "").strip()
        if not name:
            unmapped.append(sfzh)
            continue
        sql = f"UPDATE yggl SET `{col}` = %s WHERE name = %s"
        n = db.execute_update(sql, (val if val else None, name))
        if n > 0:
            updated += n
    return {
        "success": True,
        "updated": updated,
        "unmapped": unmapped,
        "message": f"已更新 {updated} 条；未匹配到 {len(unmapped)} 个身份证号",
    }
