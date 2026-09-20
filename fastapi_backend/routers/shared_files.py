# -*- coding: utf-8 -*-
"""Shared documents: list/folders/upload/rename/delete + ONLYOFFICE editor-config."""
from __future__ import annotations

import logging
import mimetypes
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple
from urllib.parse import quote

from fastapi import APIRouter, File, Form, HTTPException, Query, Request, UploadFile
from fastapi.responses import FileResponse
from openpyxl import Workbook
from pydantic import BaseModel, Field

from config import settings
from database import db
from services import onlyoffice_security as oo_sec
from services import shared_file_storage as storage

logger = logging.getLogger(__name__)

router = APIRouter(prefix="/shared-files", tags=["shared-files"])

OTHER_DEPT_LSYS = ("其他部门员工", "其他部门成员")

ALLOWED_UPLOAD_EXT = {"doc", "docx", "xls", "xlsx", "ppt", "pptx"}

MIME_BY_EXT = {
    "doc": "application/msword",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xls": "application/vnd.ms-excel",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "ppt": "application/vnd.ms-powerpoint",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}

DOC_TYPE_BY_EXT = {
    "doc": "word",
    "docx": "word",
    "xls": "cell",
    "xlsx": "cell",
    "ppt": "slide",
    "pptx": "slide",
}

DDL = """
CREATE TABLE IF NOT EXISTS shared_files (
  id BIGINT AUTO_INCREMENT PRIMARY KEY,
  name VARCHAR(500) NOT NULL,
  file_type VARCHAR(50) NOT NULL DEFAULT '',
  mime_type VARCHAR(200) NOT NULL DEFAULT '',
  parent_id BIGINT NULL,
  storage_path VARCHAR(1000) NOT NULL DEFAULT '',
  size BIGINT NOT NULL DEFAULT 0,
  owner_id VARCHAR(100) NOT NULL DEFAULT '',
  created_by VARCHAR(100) NOT NULL DEFAULT '',
  updated_by VARCHAR(100) NOT NULL DEFAULT '',
  version INT NOT NULL DEFAULT 1,
  is_folder TINYINT NOT NULL DEFAULT 0,
  is_deleted TINYINT NOT NULL DEFAULT 0,
  last_saved_key VARCHAR(100) NOT NULL DEFAULT '',
  last_saved_url VARCHAR(2000) NOT NULL DEFAULT '',
  created_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP,
  updated_at DATETIME NOT NULL DEFAULT CURRENT_TIMESTAMP ON UPDATE CURRENT_TIMESTAMP,
  KEY idx_parent_deleted (parent_id, is_deleted),
  KEY idx_owner (owner_id),
  KEY idx_deleted (is_deleted)
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COMMENT='shared documents'
"""

_NAME_BAD = re.compile(r'[\\/:*?"<>|\x00-\x1f]')


def _ensure_tables() -> None:
    try:
        db.execute_update(DDL, ())
    except Exception as e:
        logger.error("create shared_files failed: %s", e)


def _get_admin1() -> Optional[str]:
    try:
        from routers.db_manager import _get_admin1 as _admin1
        return _admin1()
    except Exception:
        return None


def _get_user_row(name: str) -> Optional[Dict[str, Any]]:
    clean = (name or "").strip()
    if not clean:
        return None
    rows = db.execute_query(
        "SELECT name, lsys, jb, gh FROM yggl WHERE name=%s AND COALESCE(zaizhi,0)=0 LIMIT 1",
        (clean,),
    )
    return rows[0] if rows else None


def is_admin1(name: str) -> bool:
    admin1 = _get_admin1()
    return bool(admin1 and (name or "").strip() == admin1)


def can_access_shared_docs(name: str) -> Tuple[bool, str]:
    user = _get_user_row(name)
    if not user:
        return False, "用户不存在或已离职"
    if is_admin1(name):
        return True, ""
    lsys = (user.get("lsys") or "").strip()
    if lsys in OTHER_DEPT_LSYS:
        return False, "外部门用户暂不可访问共享文档"
    if not lsys:
        return False, "未分配科室，无法访问共享文档"
    return True, ""


def require_module_access(name: str) -> Dict[str, Any]:
    ok, msg = can_access_shared_docs(name)
    if not ok:
        logger.warning("shared-docs denied user=%s reason=%s", name, msg)
        raise HTTPException(status_code=403, detail=msg or "没有权限")
    user = _get_user_row(name)
    assert user is not None
    return user


def can_mutate_meta(user_name: str, row: Dict[str, Any]) -> bool:
    """Rename/delete: admin1 or creator."""
    if is_admin1(user_name):
        return True
    creator = (row.get("created_by") or "").strip()
    return bool(creator and creator == (user_name or "").strip())


def get_file_row(file_id: int, *, include_deleted: bool = False) -> Optional[Dict[str, Any]]:
    if include_deleted:
        rows = db.execute_query("SELECT * FROM shared_files WHERE id=%s LIMIT 1", (file_id,))
    else:
        rows = db.execute_query(
            "SELECT * FROM shared_files WHERE id=%s AND COALESCE(is_deleted,0)=0 LIMIT 1",
            (file_id,),
        )
    return rows[0] if rows else None


def _now() -> str:
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


def _clean_name(name: str) -> str:
    clean = (name or "").strip()
    if not clean:
        raise HTTPException(status_code=400, detail="名称不能为空")
    if _NAME_BAD.search(clean) or clean in (".", ".."):
        raise HTTPException(status_code=400, detail="名称包含非法字符")
    if len(clean) > 200:
        clean = clean[:200]
    return clean


def _normalize_parent_id(parent_id: Optional[int]) -> Optional[int]:
    if parent_id is None or int(parent_id) <= 0:
        return None
    return int(parent_id)


def _assert_parent_folder(parent_id: Optional[int]) -> None:
    if parent_id is None:
        return
    parent = get_file_row(parent_id)
    if not parent:
        raise HTTPException(status_code=404, detail="父目录不存在")
    if not parent.get("is_folder"):
        raise HTTPException(status_code=400, detail="父级必须是文件夹")


def _name_exists(parent_id: Optional[int], name: str, *, exclude_id: Optional[int] = None) -> bool:
    if parent_id is None:
        sql = (
            "SELECT id FROM shared_files WHERE parent_id IS NULL AND name=%s "
            "AND COALESCE(is_deleted,0)=0"
        )
        params: tuple = (name,)
    else:
        sql = (
            "SELECT id FROM shared_files WHERE parent_id=%s AND name=%s "
            "AND COALESCE(is_deleted,0)=0"
        )
        params = (parent_id, name)
    if exclude_id:
        sql += " AND id<>%s"
        params = params + (exclude_id,)
    sql += " LIMIT 1"
    return bool(db.execute_query(sql, params))


def _max_upload_bytes() -> int:
    mb = int(getattr(settings, "SHARED_FILES_MAX_SIZE_MB", 50) or 50)
    return max(1, mb) * 1024 * 1024


def _serialize_item(row: Dict[str, Any], current_user: str) -> Dict[str, Any]:
    fid = int(row["id"])
    version = int(row.get("version") or 1)
    is_folder = bool(row.get("is_folder"))
    file_type = (row.get("file_type") or "").lower()
    editable = (not is_folder) and file_type in DOC_TYPE_BY_EXT
    return {
        "id": fid,
        "name": row.get("name"),
        "file_type": file_type,
        "mime_type": row.get("mime_type") or "",
        "parent_id": row.get("parent_id"),
        "size": int(row.get("size") or 0),
        "version": version,
        "is_folder": is_folder,
        "created_by": row.get("created_by") or "",
        "updated_by": row.get("updated_by") or "",
        "created_at": str(row.get("created_at") or ""),
        "updated_at": str(row.get("updated_at") or ""),
        "document_key": None if is_folder else oo_sec.build_document_key(fid, version),
        "can_edit": editable,
        "can_rename": can_mutate_meta(current_user, row),
        "can_delete": can_mutate_meta(current_user, row),
    }


def _breadcrumb(file_id: Optional[int]) -> List[Dict[str, Any]]:
    crumbs: List[Dict[str, Any]] = []
    seen = set()
    cur = file_id
    while cur:
        if cur in seen:
            break
        seen.add(cur)
        row = get_file_row(cur)
        if not row:
            break
        crumbs.append({"id": int(row["id"]), "name": row.get("name") or ""})
        parent = row.get("parent_id")
        cur = int(parent) if parent else None
    crumbs.reverse()
    return crumbs


def _ensure_phase2_test_file() -> None:
    try:
        existing = db.execute_query(
            "SELECT id, storage_path FROM shared_files "
            "WHERE name=%s AND COALESCE(is_folder,0)=0 AND COALESCE(is_deleted,0)=0 "
            "ORDER BY id ASC LIMIT 1",
            ("test.xlsx",),
        )
        if existing:
            row = existing[0]
            try:
                path = storage.resolve_storage_path(row.get("storage_path") or "")
                if path.is_file():
                    return
                logger.warning("test.xlsx meta exists but disk missing, rebuild id=%s", row.get("id"))
            except Exception:
                logger.warning("test.xlsx storage_path invalid, rebuild id=%s", row.get("id"))

        rel = storage.build_storage_relpath("xlsx")
        abs_path = storage.resolve_storage_path(rel)
        abs_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        ws["A1"] = "ONLYOFFICE Phase2 Test"
        ws["B1"] = "Edit this cell, close editor, reopen to verify save"
        ws["A2"] = "Phase2/Phase3 collaboration"
        wb.save(str(abs_path))
        size = abs_path.stat().st_size
        now = _now()

        if existing:
            db.execute_update(
                "UPDATE shared_files SET storage_path=%s, size=%s, version=1, "
                "mime_type=%s, file_type=%s, updated_at=%s, last_saved_key='', last_saved_url='', "
                "is_deleted=0 WHERE id=%s",
                (rel, size, MIME_BY_EXT["xlsx"], "xlsx", now, existing[0]["id"]),
            )
            logger.info("rebuilt Phase2 test.xlsx id=%s path=%s", existing[0]["id"], rel)
            return

        new_id = db.execute_insert(
            "INSERT INTO shared_files "
            "(name, file_type, mime_type, parent_id, storage_path, size, owner_id, "
            " created_by, updated_by, version, is_folder, is_deleted, "
            " last_saved_key, last_saved_url, created_at, updated_at) "
            "VALUES (%s,%s,%s,NULL,%s,%s,%s,%s,%s,1,0,0,'','',%s,%s)",
            (
                "test.xlsx",
                "xlsx",
                MIME_BY_EXT["xlsx"],
                rel,
                size,
                "system",
                "system",
                "system",
                now,
                now,
            ),
        )
        logger.info("created Phase2 test.xlsx id=%s path=%s", new_id, rel)
    except Exception as e:
        logger.error("init test.xlsx failed: %s", e)


_ensure_tables()
_ensure_phase2_test_file()


def _public_download_url(file_id: int, token: str) -> str:
    base = oo_sec.get_public_api_base()
    return f"{base}/api/shared-files/{int(file_id)}/download?token={quote(token, safe='')}"


def _public_callback_url(file_id: int) -> str:
    base = oo_sec.get_public_api_base()
    return f"{base}/api/onlyoffice/callback/{int(file_id)}"


def _editor_user_payload(user: Dict[str, Any]) -> Dict[str, str]:
    name = (user.get("name") or "").strip()
    gh = (user.get("gh") or "").strip()
    # Stable id helps ONLYOFFICE show the same collaborator across reloads.
    return {"id": gh or name, "name": name}


class FolderCreateRequest(BaseModel):
    name: str
    parent_id: Optional[int] = None
    current_user: str


class RenameRequest(BaseModel):
    name: str
    current_user: str


@router.get("")
@router.get("/")
def list_shared_files(
    current_user: str = Query(...),
    parent_id: Optional[int] = Query(None),
):
    require_module_access(current_user)
    pid = _normalize_parent_id(parent_id)
    if pid is not None:
        _assert_parent_folder(pid)
        rows = db.execute_query(
            "SELECT * FROM shared_files WHERE parent_id=%s AND COALESCE(is_deleted,0)=0 "
            "ORDER BY is_folder DESC, name ASC, id ASC",
            (pid,),
        )
    else:
        rows = db.execute_query(
            "SELECT * FROM shared_files WHERE parent_id IS NULL AND COALESCE(is_deleted,0)=0 "
            "ORDER BY is_folder DESC, name ASC, id ASC",
            (),
        )
    items = [_serialize_item(r, current_user) for r in (rows or [])]
    return {
        "success": True,
        "parent_id": pid,
        "breadcrumb": _breadcrumb(pid),
        "items": items,
        "permissions": {
            "can_upload": True,
            "can_create_folder": True,
            "is_admin": is_admin1(current_user),
        },
    }


@router.get("/phase2-test-info")
def phase2_test_info(current_user: str = Query(...)):
    require_module_access(current_user)
    rows = db.execute_query(
        "SELECT id, name, version, size, updated_at, storage_path FROM shared_files "
        "WHERE name=%s AND COALESCE(is_deleted,0)=0 ORDER BY id ASC LIMIT 1",
        ("test.xlsx",),
    )
    if not rows:
        raise HTTPException(status_code=404, detail="test.xlsx not found")
    r = rows[0]
    return {
        "success": True,
        "data": {
            "id": r["id"],
            "name": r.get("name"),
            "version": r.get("version"),
            "size": r.get("size"),
            "updated_at": str(r.get("updated_at") or ""),
            "edit_path": f"/shared-files/edit/{r['id']}",
        },
    }


@router.post("/folders")
def create_folder(req: FolderCreateRequest):
    user = require_module_access(req.current_user)
    name = _clean_name(req.name)
    pid = _normalize_parent_id(req.parent_id)
    _assert_parent_folder(pid)
    if _name_exists(pid, name):
        raise HTTPException(status_code=400, detail="同目录下已存在同名项目")
    uname = (user.get("name") or "").strip()
    now = _now()
    new_id = db.execute_insert(
        "INSERT INTO shared_files "
        "(name, file_type, mime_type, parent_id, storage_path, size, owner_id, "
        " created_by, updated_by, version, is_folder, is_deleted, "
        " last_saved_key, last_saved_url, created_at, updated_at) "
        "VALUES (%s,'folder','',%s,'',0,%s,%s,%s,1,1,0,'','',%s,%s)",
        (name, pid, uname, uname, uname, now, now),
    )
    if not new_id:
        raise HTTPException(status_code=500, detail="create folder failed")
    logger.info("shared folder create id=%s user=%s parent=%s name=%s", new_id, uname, pid, name)
    row = get_file_row(int(new_id))
    return {"success": True, "message": "文件夹已创建", "data": _serialize_item(row, uname)}


@router.post("/upload")
async def upload_file(
    current_user: str = Form(...),
    parent_id: Optional[int] = Form(None),
    file: UploadFile = File(...),
):
    user = require_module_access(current_user)
    uname = (user.get("name") or "").strip()
    pid = _normalize_parent_id(parent_id)
    _assert_parent_folder(pid)

    raw_name = (file.filename or "").strip()
    if not raw_name:
        raise HTTPException(status_code=400, detail="名称不能为空")
    # Use basename only; never trust path segments from client.
    raw_name = os.path.basename(raw_name.replace("\\", "/"))
    ext = Path(raw_name).suffix.lower().lstrip(".")
    if ext not in ALLOWED_UPLOAD_EXT:
        raise HTTPException(status_code=400, detail="仅支持 Word/Excel/PPT 文件")
    display_name = _clean_name(raw_name)
    if _name_exists(pid, display_name):
        raise HTTPException(status_code=400, detail="同目录下已存在同名项目")

    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="文件内容为空")
    if len(content) > _max_upload_bytes():
        raise HTTPException(status_code=400, detail="文件超过大小限制")

    rel = storage.build_storage_relpath(ext)
    abs_path = storage.resolve_storage_path(rel)
    size = storage.write_bytes_atomic(abs_path, content)
    now = _now()
    mime = MIME_BY_EXT.get(ext) or (file.content_type or "")
    new_id = db.execute_insert(
        "INSERT INTO shared_files "
        "(name, file_type, mime_type, parent_id, storage_path, size, owner_id, "
        " created_by, updated_by, version, is_folder, is_deleted, "
        " last_saved_key, last_saved_url, created_at, updated_at) "
        "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,1,0,0,'','',%s,%s)",
        (display_name, ext, mime, pid, rel, size, uname, uname, uname, now, now),
    )
    if not new_id:
        try:
            abs_path.unlink(missing_ok=True)
        except Exception:
            pass
        raise HTTPException(status_code=500, detail="upload failed")
    logger.info("shared upload id=%s user=%s name=%s size=%s", new_id, uname, display_name, size)
    row = get_file_row(int(new_id))
    return {"success": True, "message": "上传成功", "data": _serialize_item(row, uname)}


@router.get("/{file_id}")
def get_file_meta(file_id: int, current_user: str = Query(...)):
    require_module_access(current_user)
    row = get_file_row(file_id)
    if not row:
        raise HTTPException(status_code=404, detail="文件不存在或已删除")
    return {
        "success": True,
        "data": _serialize_item(row, current_user),
        "breadcrumb": _breadcrumb(file_id if row.get("is_folder") else row.get("parent_id")),
    }


@router.get("/{file_id}/download")
def download_file(
    file_id: int,
    request: Request,
    current_user: str = Query(""),
    token: str = Query(""),
):
    row = get_file_row(file_id)
    if not row:
        logger.warning("download missing file_id=%s", file_id)
        raise HTTPException(status_code=404, detail="文件不存在或已删除")
    if row.get("is_folder"):
        raise HTTPException(status_code=400, detail="文件夹不可下载")

    token_clean = (token or "").strip()
    user_clean = (current_user or "").strip()

    if token_clean:
        try:
            oo_sec.verify_onlyoffice_download_token(token_clean, expected_file_id=file_id)
        except Exception:
            logger.warning("onlyoffice download token rejected file_id=%s", file_id)
            raise HTTPException(status_code=403, detail="下载凭证无效或已过期")
        logger.info("shared download(onlyoffice) file_id=%s ok", file_id)
    elif user_clean:
        require_module_access(user_clean)
        logger.info("shared download(browser) file_id=%s user=%s ok", file_id, user_clean)
    else:
        raise HTTPException(status_code=401, detail="缺少下载凭证")

    try:
        path = storage.resolve_storage_path(row.get("storage_path") or "")
    except ValueError:
        raise HTTPException(status_code=500, detail="文件路径无效")
    if not path.is_file():
        logger.error("disk missing file_id=%s path=%s", file_id, path)
        raise HTTPException(status_code=404, detail="磁盘文件不存在")

    filename = (row.get("name") or path.name).strip() or path.name
    media = (row.get("mime_type") or "").strip() or mimetypes.guess_type(filename)[0] or "application/octet-stream"
    return FileResponse(
        path=str(path),
        media_type=media,
        filename=filename,
        content_disposition_type="attachment",
    )


@router.patch("/{file_id}")
def rename_file(file_id: int, req: RenameRequest):
    user = require_module_access(req.current_user)
    uname = (user.get("name") or "").strip()
    row = get_file_row(file_id)
    if not row:
        raise HTTPException(status_code=404, detail="文件不存在或已删除")
    if not can_mutate_meta(uname, row):
        logger.warning("rename denied file_id=%s user=%s", file_id, uname)
        raise HTTPException(status_code=403, detail="仅管理员或创建人可重命名")

    new_name = _clean_name(req.name)
    # Keep extension for files if user strips it accidentally.
    if not row.get("is_folder"):
        old_ext = Path(row.get("name") or "").suffix
        new_ext = Path(new_name).suffix
        if old_ext and not new_ext:
            new_name = new_name + old_ext

    parent_id = row.get("parent_id")
    pid = int(parent_id) if parent_id else None
    if _name_exists(pid, new_name, exclude_id=file_id):
        raise HTTPException(status_code=400, detail="同目录下已存在同名项目")

    now = _now()
    affected = db.execute_update(
        "UPDATE shared_files SET name=%s, updated_by=%s, updated_at=%s "
        "WHERE id=%s AND COALESCE(is_deleted,0)=0",
        (new_name, uname, now, file_id),
    )
    if not affected:
        raise HTTPException(status_code=500, detail="rename failed")
    logger.info("shared rename file_id=%s user=%s name=%s", file_id, uname, new_name)
    fresh = get_file_row(file_id)
    return {"success": True, "message": "已重命名", "data": _serialize_item(fresh, uname)}


@router.delete("/{file_id}")
def delete_file(file_id: int, current_user: str = Query(...)):
    user = require_module_access(current_user)
    uname = (user.get("name") or "").strip()
    row = get_file_row(file_id)
    if not row:
        raise HTTPException(status_code=404, detail="文件不存在或已删除")
    if not can_mutate_meta(uname, row):
        logger.warning("delete denied file_id=%s user=%s", file_id, uname)
        raise HTTPException(status_code=403, detail="仅管理员或创建人可删除")

    now = _now()
    # Soft-delete node + all descendants (folder tree).
    to_delete = [file_id]
    queue = [file_id]
    while queue:
        cur = queue.pop(0)
        kids = db.execute_query(
            "SELECT id FROM shared_files WHERE parent_id=%s AND COALESCE(is_deleted,0)=0",
            (cur,),
        )
        for k in kids or []:
            kid = int(k["id"])
            to_delete.append(kid)
            queue.append(kid)

    for fid in to_delete:
        db.execute_update(
            "UPDATE shared_files SET is_deleted=1, updated_by=%s, updated_at=%s WHERE id=%s",
            (uname, now, fid),
        )
    logger.info("shared soft-delete file_id=%s user=%s count=%s", file_id, uname, len(to_delete))
    return {"success": True, "message": "已删除", "deleted_count": len(to_delete)}


@router.get("/{file_id}/editor-config")
def get_editor_config(file_id: int, current_user: str = Query(...)):
    user = require_module_access(current_user)
    row = get_file_row(file_id)
    if not row:
        raise HTTPException(status_code=404, detail="文件不存在或已删除")
    if row.get("is_folder"):
        raise HTTPException(status_code=400, detail="文件夹无法在线编辑")

    ext = (row.get("file_type") or Path(row.get("name") or "").suffix.lstrip(".")).lower()
    if ext not in DOC_TYPE_BY_EXT:
        raise HTTPException(status_code=400, detail=f"不支持的文件类型: {ext}")

    try:
        path = storage.resolve_storage_path(row.get("storage_path") or "")
    except ValueError:
        raise HTTPException(status_code=500, detail="文件路径无效")
    if not path.is_file():
        raise HTTPException(status_code=404, detail="磁盘文件不存在")

    try:
        document_server_url = oo_sec.get_onlyoffice_url()
        version = int(row.get("version") or 1)
        doc_key = oo_sec.build_document_key(file_id, version)
        download_token = oo_sec.issue_onlyoffice_download_token(
            file_id=file_id,
            version=version,
            document_key=doc_key,
        )
        user_payload = _editor_user_payload(user)
        config = {
            "document": {
                "fileType": ext,
                "key": doc_key,
                "title": row.get("name") or f"file.{ext}",
                "url": _public_download_url(file_id, download_token),
                "permissions": {
                    "edit": True,
                    "download": True,
                    "print": True,
                    "review": True,
                    "comment": True,
                },
            },
            "documentType": DOC_TYPE_BY_EXT[ext],
            "editorConfig": {
                "callbackUrl": _public_callback_url(file_id),
                "mode": "edit",
                "lang": "zh-CN",
                "user": user_payload,
                "customization": {
                    "forcesave": True,
                    "autosave": True,
                },
            },
        }
        config["token"] = oo_sec.sign_editor_config(config)
    except oo_sec.OnlyOfficeConfigError as e:
        logger.error("editor-config config error file_id=%s: %s", file_id, e)
        raise HTTPException(status_code=503, detail=str(e))
    except Exception as e:
        logger.exception("editor-config failed file_id=%s", file_id)
        raise HTTPException(status_code=500, detail=f"生成编辑配置失败: {e}")

    logger.info(
        "open editor file_id=%s user=%s user_id=%s key=%s version=%s",
        file_id,
        current_user,
        user_payload.get("id"),
        config["document"]["key"],
        version,
    )
    return {
        "success": True,
        "documentServerUrl": document_server_url,
        "config": config,
        "file": {
            "id": file_id,
            "name": row.get("name"),
            "version": version,
            "document_key": doc_key,
        },
    }
