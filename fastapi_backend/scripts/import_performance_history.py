# -*- coding: utf-8 -*-
"""一次性导入旧版月度绩效工作簿。

示例：
  python scripts/import_performance_history.py "../2025年员工绩效考核 (1).xlsx" --year 2025 --apply
不带 --apply 仅打印将导入的记录数，不会写入数据库。
"""
import argparse
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from database import db, db_demo  # noqa: E402


def read_sheet(ws, default_year):
    match = re.match(r"(?:(\d{2})\.)?(\d{1,2})月绩效$", ws.title.strip())
    if not match:
        return []
    year = 2000 + int(match.group(1)) if match.group(1) else default_year
    month = int(match.group(2))
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    header = {str(value).strip(): index for index, value in enumerate(rows[0]) if value is not None}
    name_idx = header.get("姓名", header.get("12月总排名"))
    score_idx = header.get("绩效得分")
    if name_idx is None or score_idx is None:
        return []
    dept_idx, title_idx = header.get("班组"), header.get("职称", header.get("津贴"))
    rank_idx, pct_idx = header.get("排名"), header.get("排名百分比")
    result = []
    for row in rows[1:]:
        name = str(row[name_idx] or "").strip() if name_idx < len(row) else ""
        score = row[score_idx] if score_idx < len(row) else None
        if not name or score is None:
            continue
        try:
            score = float(score)
        except (TypeError, ValueError):
            continue
        dept = str(row[dept_idx] or "未分组").strip() if dept_idx is not None and dept_idx < len(row) else "未分组"
        title = str(row[title_idx] or "").strip() if title_idx is not None and title_idx < len(row) else ""
        rank = row[rank_idx] if rank_idx is not None and rank_idx < len(row) else None
        pct = row[pct_idx] if pct_idx is not None and pct_idx < len(row) else None
        result.append((date(year, month, 1), dept, name, title, score, rank, pct))
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--year", type=int, required=True, help="没有年份前缀的月份页所属年份")
    parser.add_argument("--apply", action="store_true", help="确认写入数据库")
    args = parser.parse_args()
    book = openpyxl.load_workbook(args.workbook, read_only=True, data_only=True)
    records = [record for sheet in book.worksheets for record in read_sheet(sheet, args.year)]
    names = sorted({record[2] for record in records})
    levels = {}
    if names:
        placeholders = ",".join(["%s"] * len(names))
        source_levels = db_demo.execute_query(
            f"SELECT name, job_level FROM employee_info WHERE name IN ({placeholders})", tuple(names)
        )
        levels = {str(row.get("name") or "").strip(): row.get("job_level") or "" for row in source_levels}
    current_depts = {}
    if names:
        placeholders = ",".join(["%s"] * len(names))
        source_depts = db.execute_query(
            f"SELECT name, lsys FROM yggl WHERE name IN ({placeholders})", tuple(names)
        )
        current_depts = {str(row.get("name") or "").strip(): (row.get("lsys") or "").strip() for row in source_depts}
    # 新系统统一以 demo.employee_info.job_level 为准；历史班组尽量映射到当前 yggl.lsys，
    # 以确保现任班组长在“统计汇总”可直接查到本组旧数据。
    records = [
        (record[0], current_depts.get(record[2]) or record[1], record[2], levels.get(record[2]) or record[3], *record[4:])
        for record in records
    ]
    print(f"发现 {len(records)} 条历史绩效记录")
    if not args.apply:
        print("预览完成；添加 --apply 后写入数据库。")
        return
    sql = """
      INSERT INTO performance_records
        (performance_month, department, employee_name, job_level, score, rank_no, rank_percent, created_by)
      VALUES (%s,%s,%s,%s,%s,%s,%s,'历史Excel导入')
      ON DUPLICATE KEY UPDATE department=VALUES(department), job_level=VALUES(job_level), score=VALUES(score),
        rank_no=VALUES(rank_no), rank_percent=VALUES(rank_percent), created_by=VALUES(created_by)
    """
    affected = db.execute_many(sql, records)
    if affected < 0:
        raise SystemExit("导入失败")
    print(f"已导入/更新 {len(records)} 条历史绩效记录。")


if __name__ == "__main__":
    main()
